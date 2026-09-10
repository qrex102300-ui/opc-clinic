#!/usr/bin/env python3
"""Synthetic RFQ-to-quote proof.

Phase-1 constraints:
- text-based PDF with one agreed line format
- exact product-code and unit matching only
- prices only from an approved price list
- no substitutions
- uncertain rows go to human review
- Excel draft is never treated as approved
"""
from __future__ import annotations

import argparse
import csv
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from pypdf import PdfReader

LINE_RE = re.compile(
    r"^\s*(?P<code>[^|]+?)\s*\|\s*(?P<description>[^|]*?)\s*\|\s*"
    r"(?P<quantity>[^|]+?)\s*\|\s*(?P<unit>[^|]+?)\s*$"
)


def norm(value: str) -> str:
    return (value or "").strip().upper()


def read_csv(path: str | Path) -> list[dict[str, str]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def extract_text_pdf(path: str | Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def parse_rfq_text(text: str) -> list[dict[str, str]]:
    rows = []
    for raw in text.splitlines():
        match = LINE_RE.match(raw)
        if not match:
            continue
        row = {k: v.strip() for k, v in match.groupdict().items()}
        if norm(row["code"]) in {"CODE", "PRODUCT CODE", "SKU"}:
            continue
        rows.append(row)
    return rows


def _index(rows: list[dict[str, str]], *fields: str):
    out: dict[tuple[str, ...], list[dict[str, str]]] = {}
    for row in rows:
        key = tuple(norm(row.get(f, "")) for f in fields)
        out.setdefault(key, []).append(row)
    return out


def prepare_quote(rfq_rows, catalogue_rows, price_rows):
    catalog_by_code = _index(catalogue_rows, "code")
    price_by_code_unit = _index(price_rows, "code", "unit")
    results = []

    for i, src in enumerate(rfq_rows, start=1):
        code = norm(src.get("code", ""))
        unit = norm(src.get("unit", ""))
        reasons = []
        approved_price_minor = ""
        line_total_minor = ""
        price_list_version = ""
        currency = ""

        try:
            quantity = Decimal(str(src.get("quantity", "")).strip())
            if quantity <= 0 or quantity != quantity.to_integral_value():
                raise InvalidOperation
        except (InvalidOperation, ValueError):
            quantity = None
            reasons.append("INVALID_QUANTITY")

        catalog_matches = catalog_by_code.get((code,), [])
        if not code:
            reasons.append("MISSING_CODE")
        elif len(catalog_matches) == 0:
            reasons.append("UNKNOWN_CODE")
        elif len(catalog_matches) > 1:
            reasons.append("AMBIGUOUS_CATALOG_CODE")
        else:
            approved_unit = norm(catalog_matches[0].get("unit", ""))
            if unit != approved_unit:
                reasons.append("UNIT_MISMATCH")

        if not reasons:
            price_matches = price_by_code_unit.get((code, unit), [])
            if len(price_matches) == 0:
                reasons.append("MISSING_APPROVED_PRICE")
            elif len(price_matches) > 1:
                reasons.append("AMBIGUOUS_APPROVED_PRICE")
            else:
                price = price_matches[0]
                try:
                    minor = int(str(price.get("price_minor", "")).strip())
                    if minor < 0:
                        raise ValueError
                except ValueError:
                    reasons.append("INVALID_APPROVED_PRICE")
                else:
                    approved_price_minor = str(minor)
                    line_total_minor = str(minor * int(quantity))
                    price_list_version = str(price.get("price_list_version", "")).strip()
                    currency = norm(price.get("currency", ""))
                    if not price_list_version or not currency:
                        reasons.append("PRICE_PROVENANCE_MISSING")
                        approved_price_minor = ""
                        line_total_minor = ""

        results.append({
            "line": str(i),
            "code": src.get("code", "").strip(),
            "description": src.get("description", "").strip(),
            "quantity": src.get("quantity", "").strip(),
            "unit": src.get("unit", "").strip(),
            "status": "REVIEW" if reasons else "OK",
            "reason": ";".join(reasons),
            "approved_price_minor": approved_price_minor,
            "line_total_minor": line_total_minor,
            "currency": currency,
            "price_list_version": price_list_version,
        })
    return results


def write_excel(results, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"
    headers = ["line", "code", "description", "quantity", "unit", "status", "reason",
               "approved_price_minor", "line_total_minor", "currency", "price_list_version"]
    ws.append(headers)
    for row in results:
        ws.append([row[h] for h in headers])

    exc = wb.create_sheet("Exceptions")
    exc.append(headers)
    for row in results:
        if row["status"] != "OK":
            exc.append([row[h] for h in headers])

    summary = wb.create_sheet("Summary")
    summary.append(["field", "value"])
    blocked = any(r["status"] != "OK" for r in results)
    summary.append(["draft_status", "REVIEW_REQUIRED" if blocked else "DRAFT_READY"])
    summary.append(["quote_total_minor", "WITHHELD" if blocked else str(sum(int(r["line_total_minor"]) for r in results))])
    versions = sorted({r["price_list_version"] for r in results if r["price_list_version"]})
    summary.append(["price_list_versions", ",".join(versions)])
    summary.append(["approval", "HUMAN_APPROVAL_REQUIRED"])
    wb.save(str(output_path))


def run(pdf_path, catalogue_csv, price_csv, output_xlsx):
    rfq = parse_rfq_text(extract_text_pdf(pdf_path))
    if not rfq:
        raise ValueError("No RFQ rows matched the agreed text layout")
    results = prepare_quote(rfq, read_csv(catalogue_csv), read_csv(price_csv))
    write_excel(results, output_xlsx)
    return results


def main():
    p = argparse.ArgumentParser()
    p.add_argument("pdf")
    p.add_argument("catalogue_csv")
    p.add_argument("price_csv")
    p.add_argument("output_xlsx")
    args = p.parse_args()
    rows = run(args.pdf, args.catalogue_csv, args.price_csv, args.output_xlsx)
    ok = sum(r["status"] == "OK" for r in rows)
    print(f"processed={len(rows)} ok={ok} review={len(rows)-ok} output={args.output_xlsx}")


if __name__ == "__main__":
    main()
