import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from reportlab.pdfgen import canvas

from rfq_quote import run


class RFQQuoteDemoTest(unittest.TestCase):
    def _csv(self, path, fields, rows):
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)

    def test_pdf_to_excel_and_review_gates(self):
        with tempfile.TemporaryDirectory() as td:
            td = Path(td)
            pdf = td / "rfq.pdf"
            c = canvas.Canvas(str(pdf))
            y = 800
            for line in [
                "CODE | DESCRIPTION | QTY | UNIT",
                "SW-01 | Main switch | 3 | EA",
                "ZZ-99 | Unknown item | 2 | EA",
                "SW-01 | Wrong unit | 1 | BOX",
                "SW-02 | Known but unpriced | 4 | EA",
            ]:
                c.drawString(72, y, line)
                y -= 24
            c.save()

            catalogue = td / "catalogue.csv"
            prices = td / "prices.csv"
            output = td / "quote.xlsx"
            self._csv(catalogue, ["code", "description", "unit"], [
                {"code": "SW-01", "description": "Synthetic switch", "unit": "EA"},
                {"code": "SW-02", "description": "Synthetic relay", "unit": "EA"},
            ])
            self._csv(prices, ["code", "unit", "price_minor", "currency", "price_list_version"], [
                {"code": "SW-01", "unit": "EA", "price_minor": "1250", "currency": "EUR", "price_list_version": "PL-2026-09"},
            ])

            results = run(pdf, catalogue, prices, output)
            self.assertEqual(4, len(results))
            self.assertEqual("3750", results[0]["line_total_minor"])
            self.assertEqual("UNKNOWN_CODE", results[1]["reason"])
            self.assertEqual("UNIT_MISMATCH", results[2]["reason"])
            self.assertEqual("MISSING_APPROVED_PRICE", results[3]["reason"])

            wb = load_workbook(output, data_only=True)
            self.assertEqual(["Quote", "Exceptions", "Summary"], wb.sheetnames)
            self.assertEqual(5, wb["Quote"].max_row)
            self.assertEqual(4, wb["Exceptions"].max_row)
            self.assertEqual("REVIEW_REQUIRED", wb["Summary"]["B2"].value)
            self.assertEqual("WITHHELD", wb["Summary"]["B3"].value)
            self.assertEqual("HUMAN_APPROVAL_REQUIRED", wb["Summary"]["B5"].value)


if __name__ == "__main__":
    unittest.main()
